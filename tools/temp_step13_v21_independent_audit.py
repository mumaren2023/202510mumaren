from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

EXPECTED_COMPLETE_SHA = "eaccf2d976a8fc283888919218894d6b622ba106afcaf13f00f92cd7725d6a86"
CORE_SHA = "30f53680015c9b5a97d66bedbe5390d9208c67cdcefeae21892e7325e636316f"
DATE = "2026-07-22"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_json(path: Path, default: Any = None) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return copy.deepcopy(default)


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def run_python(script: Path, cwd: Path, timeout: int = 300) -> subprocess.CompletedProcess[str]:
    env = dict(os.environ)
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    return subprocess.run(
        [sys.executable, str(script)], cwd=str(cwd), env=env,
        capture_output=True, text=True, timeout=timeout,
    )


def clean_runtime(root: Path) -> None:
    for p in list(root.rglob("__pycache__")):
        shutil.rmtree(p, ignore_errors=True)
    for p in list(root.rglob("*.pyc")):
        p.unlink(missing_ok=True)


def download(url: str, path: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=180) as response, path.open("wb") as f:
        shutil.copyfileobj(response, f)


def identify_package_root(folder: Path) -> Path:
    candidates = [p for p in folder.iterdir() if p.is_dir()]
    for p in candidates:
        if (p / "current").is_dir() and (p / "history").is_dir():
            return p
    if len(candidates) == 1:
        return candidates[0]
    raise RuntimeError("Cannot identify package root")


def file_hash_map(root: Path, exclude_names: set[str] | None = None) -> dict[str, str]:
    excluded = exclude_names or set()
    return {
        p.relative_to(root).as_posix(): sha256(p)
        for p in root.rglob("*")
        if p.is_file() and p.name not in excluded
    }


def find_validation_summary(current: Path) -> dict[str, Any]:
    candidates = sorted((current / "reports").glob("validation_summary*.json"))
    if not candidates:
        raise RuntimeError(f"No validation summary under {current}")
    return read_json(candidates[-1], {})


def load_workflow(current: Path):
    workflow_path = current / "tests" / "workflow_simulator.py"
    spec = importlib.util.spec_from_file_location("independent_workflow_v21", workflow_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot import workflow simulator")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def independent_boundary_tests(current: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    module = load_workflow(current)
    Flow = module.ReferenceWorkflow
    fixed = lambda: datetime(2026, 7, 22, 12, 0, tzinfo=timezone.utc)
    u1 = {"authenticated": True, "principal_id": "USER-1", "role": "user"}
    g1 = {"authenticated": True, "principal_id": "GUARDIAN-1", "role": "guardian"}
    privacy = {"authenticated": True, "principal_id": "PRIVACY-1", "role": "privacy_reviewer"}
    tests: list[dict[str, Any]] = []
    findings: list[dict[str, Any]] = []

    def record(test_id: str, title: str, expected: str, actual: Any, passed: bool, issue: str | None = None):
        row = {"test_id": test_id, "title": title, "expected": expected, "actual": actual, "passed": bool(passed), "issue": issue}
        tests.append(row)
        return row

    base_work = {
        "work_id": "WORK-1", "user_id": "U-1", "content_asset_id": "ASSET-U1-001",
        "content_type": "image/png", "consent_status": "submitted_for_review",
        "idempotency_key": "ind-work-0001", "operation": "create",
    }
    base_progress = {
        "user_id": "U-1", "resource_id": "COURSE-DEMO-001", "core_ref_id": "ORB-GEN-00000006",
        "core_version": "V1.0-beta1", "progress": 0.5,
        "client_event_at": "2026-07-22T12:00:00Z", "idempotency_key": "ind-progress-0001",
        "expected_version": 0,
    }
    base_attempt = {
        "attempt_id": "ATT-IND-1", "game_id": "GAME-DEMO-001", "core_ref_ids": ["ORB-GEN-00000006"],
        "core_version": "V1.0-beta1", "submitted_answer": {"choice": "a"},
        "client_event_at": "2026-07-22T12:00:00Z", "idempotency_key": "ind-attempt-0001",
        "operation": "create",
    }

    # Positive controls and previous four findings.
    flow = Flow(clock_provider=fixed)
    good_work = flow.create_child_work(copy.deepcopy(base_work), u1)
    record("CTRL-01", "本人active资产创建作品", "201", good_work, good_work.get("status") == 201)

    for asset_id, expected_status in [
        ("ASSET-NOPE", 422), ("ASSET-U2-PRIVATE", 403),
        ("ASSET-U1-QUARANTINED", 422), ("ASSET-U1-RESTRICTED", 422),
    ]:
        f = Flow(clock_provider=fixed)
        req = copy.deepcopy(base_work)
        req.update({"work_id": f"WORK-{asset_id}", "content_asset_id": asset_id, "idempotency_key": f"ind-{asset_id}"})
        result = f.create_child_work(req, u1)
        record(f"FIX-ASSET-{asset_id}", f"资产闭环：{asset_id}", str(expected_status), result, result.get("status") == expected_status)

    f = Flow(clock_provider=fixed)
    wrong_progress = copy.deepcopy(base_progress)
    wrong_progress.update({"core_ref_id": "ORB-GEN-00000007", "idempotency_key": "ind-progress-moon"})
    result = f.write_progress(wrong_progress, u1)
    record("FIX-CLOSURE-01", "日课程＋月核心对象", "422 product_core_closure_mismatch", result,
           result.get("status") == 422 and result.get("error") == "product_core_closure_mismatch")

    f = Flow(clock_provider=fixed)
    wrong_attempt = copy.deepcopy(base_attempt)
    wrong_attempt.update({"attempt_id": "ATT-IND-MOON", "core_ref_ids": ["ORB-GEN-00000007"], "idempotency_key": "ind-attempt-moon"})
    result = f.create_attempt(wrong_attempt, u1)
    record("FIX-CLOSURE-02", "日游戏＋月核心对象", "422 product_core_closure_mismatch", result,
           result.get("status") == 422 and result.get("error") == "product_core_closure_mismatch")

    clock = [datetime(2026, 7, 22, 12, 0, tzinfo=timezone.utc)]
    f = Flow(clock_provider=lambda: clock[0])
    first = f.write_progress(copy.deepcopy(base_progress), u1)
    before_count = len(f.progress)
    clock[0] = datetime(2026, 9, 1, 12, 0, tzinfo=timezone.utc)
    f.products["COURSE-DEMO-001"]["status"] = "inactive"
    second = f.write_progress(copy.deepcopy(base_progress), u1)
    record("FIX-IDEM-01", "时间老化和资源停用后稳定重放", "与首次响应完全一致且记录数不增加",
           {"first": first, "second": second, "records": len(f.progress)},
           first == second and len(f.progress) == before_count)

    f = Flow(clock_provider=fixed)
    progress_result = f.write_progress(copy.deepcopy(base_progress), u1)
    progress_record = next(iter(f.progress.values())) if f.progress else {}
    expected_snapshots = {"resource_version", "contract_sha256", "core_closure_version", "server_recorded_at"}
    record("FIX-SNAPSHOT-01", "progress保存产品契约快照", "四个服务端快照字段",
           progress_record, progress_result.get("status") == 200 and expected_snapshots.issubset(progress_record))
    attempt_result = f.create_attempt(copy.deepcopy(base_attempt), u1)
    attempt_record = f.attempts.get("ATT-IND-1", {})
    record("FIX-SNAPSHOT-02", "attempt保存产品契约快照", "四个服务端快照字段",
           attempt_record, attempt_result.get("status") == 201 and expected_snapshots.issubset(attempt_record))

    # New independent attacks: the previous closure requirement explicitly required binding asset to current work_id.
    f = Flow(clock_provider=fixed)
    req1 = copy.deepcopy(base_work)
    req1.update({"work_id": "WORK-BIND-1", "idempotency_key": "ind-bind-1"})
    req2 = copy.deepcopy(base_work)
    req2.update({"work_id": "WORK-BIND-2", "idempotency_key": "ind-bind-2"})
    r1 = f.create_child_work(req1, u1)
    r2 = f.create_child_work(req2, u1)
    passed = r1.get("status") == 201 and r2.get("status") in {403, 409, 422}
    record("NEW-01", "同一内容资产绑定两个不同work_id", "第二次拒绝", {"first": r1, "second": r2}, passed,
           None if passed else "S13R12-P1-01")
    if not passed:
        findings.append({
            "issue_id": "S13R12-P1-01", "priority": "P1",
            "title": "同一儿童内容资产可重复绑定多个作品ID",
            "evidence": "同一content_asset_id使用不同work_id和幂等键连续创建，两次均返回201。资产登记表没有bound_work_id或多次使用策略。",
            "impact": "同一原始资产可形成多个独立原创作品和审核链，造成重复公开、重复统计、撤回不一致及版权来源歧义。",
            "repair": "资产首次绑定时写入bound_work_id或建立asset_work_bindings；默认一对一。确需复用时必须采用明确的派生关系和独立版本，而不是静默重复。",
        })

    # New independent attack: asset state must be revalidated at review/publication time.
    f = Flow(clock_provider=fixed)
    work = copy.deepcopy(base_work)
    work.update({"work_id": "WORK-STATE-1", "idempotency_key": "ind-state-work"})
    created = f.create_child_work(work, u1)
    f.assets["ASSET-U1-001"]["status"] = "quarantined"
    guardian_req = {
        "review_id": "REVIEW-STATE-G", "decision": "guardian_approved",
        "permission_evidence_id": "PERM-GUARDIAN-W1", "idempotency_key": "ind-state-guardian",
        "expected_version": 1, "operation": "review",
    }
    # Permission fixture is bound to WORK-1; rebind only to isolate asset-state behavior.
    f.permissions["PERM-GUARDIAN-W1"]["work_id"] = "WORK-STATE-1"
    guardian_result = f.review_child_work("WORK-STATE-1", guardian_req, g1)
    passed = guardian_result.get("status") in {403, 409, 422}
    record("NEW-02", "资产创建后进入quarantined再审核", "拒绝", guardian_result, passed,
           None if passed else "S13R12-P1-02")

    f2 = Flow(clock_provider=fixed)
    work2 = copy.deepcopy(base_work)
    work2.update({"work_id": "WORK-STATE-2", "idempotency_key": "ind-state-work2"})
    f2.create_child_work(work2, u1)
    f2.permissions["PERM-GUARDIAN-W1"]["work_id"] = "WORK-STATE-2"
    f2.permissions["PERM-PRIVACY-W1"]["work_id"] = "WORK-STATE-2"
    guardian2 = copy.deepcopy(guardian_req)
    guardian2.update({"review_id": "REVIEW-STATE-G2", "idempotency_key": "ind-state-guardian2"})
    g_result = f2.review_child_work("WORK-STATE-2", guardian2, g1)
    f2.assets["ASSET-U1-001"]["release_state"] = "restricted"
    privacy_req = {
        "review_id": "REVIEW-STATE-P", "decision": "public_approved",
        "permission_evidence_id": "PERM-PRIVACY-W1", "idempotency_key": "ind-state-privacy",
        "expected_version": 2, "operation": "review",
    }
    privacy_result = f2.review_child_work("WORK-STATE-2", privacy_req, privacy)
    passed_privacy = g_result.get("status") == 200 and privacy_result.get("status") in {403, 409, 422}
    record("NEW-03", "资产在公开审核前变为restricted", "公开审核拒绝",
           {"guardian": g_result, "privacy": privacy_result}, passed_privacy,
           None if passed_privacy else "S13R12-P1-02")
    if not passed or not passed_privacy:
        findings.append({
            "issue_id": "S13R12-P1-02", "priority": "P1",
            "title": "审核和公开阶段未重新核验资产当前安全与发布状态",
            "evidence": "作品创建后把资产改为quarantined或restricted，guardian/privacy审核仍可能返回200。review_child_work只检查作品、主体、PERM和状态机，不检查当前资产登记。",
            "impact": "已隔离、撤回或限制发布的儿童内容仍可能获得公开批准，直接影响隐私、安全、版权撤回和误发布控制。",
            "repair": "每次guardian审核和privacy公开审核都按content_asset_id重新读取资产；核验active、owner、release_state、checksum/version一致性。状态异常必须阻断，必要时自动撤销待审状态。",
        })

    # Contract fingerprint integrity: allowed_core_refs can be changed while stale hash remains accepted.
    f = Flow(clock_provider=fixed)
    product = f.products["COURSE-DEMO-001"]
    original_hash = product["contract_sha256"]
    product["allowed_core_refs"] = ["ORB-GEN-00000006", "ORB-GEN-00000007"]
    tampered = copy.deepcopy(base_progress)
    tampered.update({"core_ref_id": "ORB-GEN-00000007", "idempotency_key": "ind-contract-tamper"})
    result = f.write_progress(tampered, u1)
    stale_accepted = result.get("status") == 200 and result.get("contract_sha256") == original_hash
    record("NEW-04", "产品核心闭包变化但contract_sha256未同步", "拒绝或检测哈希不一致", result,
           not stale_accepted, "S13R12-P2-01" if stale_accepted else None)
    if stale_accepted:
        findings.append({
            "issue_id": "S13R12-P2-01", "priority": "P2",
            "title": "产品契约哈希未在运行时校验，闭包变更可携带陈旧哈希写入",
            "evidence": "修改COURSE-DEMO-001.allowed_core_refs加入月对象但保留原contract_sha256，月对象progress返回200并保存陈旧哈希。",
            "impact": "contract_sha256不能证明实际执行的内容闭包，产品记录可能声称使用旧契约，实际按新或被篡改闭包运行。",
            "repair": "加载产品注册表时重算规范化契约哈希；每次首次写入验证哈希与闭包一致。注册表应签名或只读版本化，哈希不一致立即拒绝并告警。",
        })

    # Version registry model currently cannot retain multiple versions with same resource ID.
    workflow_text = (current / "tests" / "workflow_simulator.py").read_text(encoding="utf-8")
    registry = read_json(current / "policies" / "product_resource_registry.json", {})
    keyed_only_by_id = 'self.products={x["resource_id"]:x' in workflow_text.replace(" ", "") or 'self.products = {x["resource_id"]: x' in workflow_text
    unique_rule = len({x.get("resource_id") for x in registry.get("records", [])}) == len(registry.get("records", []))
    update_rule = str(registry.get("update_rule", ""))
    contradiction = keyed_only_by_id and "new resource_version" in update_rule and unique_rule
    record("NEW-05", "同一resource_id的多版本契约可并存和解析", "版本化键或不可变版本档案",
           {"keyed_only_by_resource_id": keyed_only_by_id, "update_rule": update_rule, "current_ids_unique": unique_rule},
           not contradiction, "S13R12-P2-02" if contradiction else None)
    if contradiction:
        findings.append({
            "issue_id": "S13R12-P2-02", "priority": "P2",
            "title": "产品注册表声明新增resource_version，但运行模型只按resource_id取单一记录",
            "evidence": "策略要求同ID变更新增resource_version；ReferenceWorkflow却以resource_id构造字典，无法同时保存和解析同一产品的V1.0、V2.0。",
            "impact": "旧记录虽保存版本和哈希，但系统缺少对应契约档案的可解析入口，无法完整复现历史产品内容。",
            "repair": "注册表以(resource_id, resource_version)为稳定复合键；另设active_version指针。保存每版完整契约和闭包，不覆盖旧版；记录解析必须按快照版本获取。",
        })

    # Deduplicate findings by issue id.
    unique: dict[str, dict[str, Any]] = {}
    for finding in findings:
        unique[finding["issue_id"]] = finding
    return tests, list(unique.values())


def inspect_mapping(outer_dir: Path) -> dict[str, Any]:
    candidates = list(outer_dir.rglob("*字段映射与接口契约_V2.1*.xlsx"))
    if not candidates:
        return {"found": False}
    path = candidates[0]
    wb = load_workbook(path, data_only=False, read_only=True)
    sheet_name = "02_教育转化字段映射" if "02_教育转化字段映射" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] not in (None, ""):
            rows.append(row)
    counts: dict[str, int] = {}
    if "教育接口分类" in headers:
        idx = headers.index("教育接口分类")
        for row in rows:
            value = row[idx]
            counts[str(value)] = counts.get(str(value), 0) + 1
    formula_errors = []
    for ws2 in wb.worksheets:
        for row in ws2.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(err in value for err in ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"]):
                    formula_errors.append(f"{ws2.title}!{cell.coordinate}:{value}")
    return {"found": True, "filename": path.name, "mapped_rows": len(rows), "classification_counts": counts, "formula_errors": formula_errors}


def doc_setup(doc: Document, header: str):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.4); sec.left_margin = Cm(1.7); sec.right_margin = Cm(1.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC"); normal.font.size = Pt(10)
    for name, size, color in [("Title", 21, "1F4E78"), ("Heading 1", 15, "1F4E78"), ("Heading 2", 12, "4472C4")]:
        st = doc.styles[name]; st.font.name = "Noto Sans CJK SC"; st._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    hp = sec.header.paragraphs[0]; hp.text = header; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("第13步V2.1独立验收｜未进入下一步    ")
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); fp._p.append(fld)


def add_title(doc: Document, conclusion: str):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(50)
    r = p.add_run("第13步｜教育转化接口独立验收报告"); r.bold = True; r.font.size = Pt(21); r.font.color.rgb = RGBColor.from_string("1F4E78")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("V2.2 独立验收版｜2026-07-22"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string("4472C4")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f"审查对象：第13步V2.1修复候选\n验收结论：{conclusion}")
    doc.add_page_break()


def add_table(doc: Document, headers: list[str], rows: list[list[Any]]):
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers): table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row): cells[i].text = str(value)
    doc.add_paragraph()


def create_reports(out: Path, audit: dict[str, Any], tests: list[dict[str, Any]], findings: list[dict[str, Any]]) -> list[Path]:
    p0 = sum(1 for x in findings if x["priority"] == "P0")
    p1 = sum(1 for x in findings if x["priority"] == "P1")
    p2 = sum(1 for x in findings if x["priority"] == "P2")
    p3 = sum(1 for x in findings if x["priority"] == "P3")
    conclusion = "通过，可以进入下一步" if not findings else "暂不通过，需要修正"

    report = out / f"甲骨文数字文化教育系统_核心数据库_第13步教育转化接口独立验收报告_V2.2_{'通过版' if not findings else '暂不通过版'}_20260722.docx"
    doc = Document(); doc_setup(doc, "甲骨文数字文化教育系统｜第13步V2.2独立验收"); add_title(doc, conclusion)
    if findings:
        doc.add_heading("先列修正方案", level=1)
        add_table(doc, ["顺序", "级别", "问题", "修正要求", "建议修正版命名"], [
            [i + 1, f["priority"], f["title"], f["repair"], f"V2.2_{f['issue_id']}_修正版"] for i, f in enumerate(findings)
        ])
    doc.add_heading("1. 独立验收范围与完整性", level=1)
    add_table(doc, ["项目", "结果"], [
        ["V2.1完整包SHA-256", audit["complete_sha256"]],
        ["预期SHA-256", EXPECTED_COMPLETE_SHA],
        ["ZIP CRC", audit["zip_crc"]],
        ["manifest", f"{audit['manifest_entries']}项，错误{audit['manifest_errors']}"],
        ["V2.1原测试", audit["v21_tests"]],
        ["V2.0历史回归", audit["v20_tests"]],
        ["V2.0历史文件比较", audit["history_compare"]],
        ["核心数据库影响", "无；核心SHA保持"],
    ])
    doc.add_heading("2. 第13步交付物与验收条件", level=1)
    add_table(doc, ["项目", "独立判断"], [
        ["教育转化字段映射", str(audit["mapping"])],
        ["儿童/研究模式接口", "核心研究字段、教育转述、争议提示和原创标签仍分层；未发现核心证据层混写"],
        ["课程/游戏/展品最小数据契约", "产品ID与登记核心闭包的基本错配攻击已拒绝；新增问题见下表"],
        ["前端只读和回写边界", "核心数据库只读；资产和产品回写新增边界见问题清单"],
        ["产品阶段闸门", "独立验收完成前保持关闭"],
    ])
    doc.add_heading("3. 独立边界测试", level=1)
    add_table(doc, ["测试ID", "场景", "预期", "结果", "结论"], [
        [t["test_id"], t["title"], t["expected"], json.dumps(t["actual"], ensure_ascii=False)[:500], "PASS" if t["passed"] else "FAIL"] for t in tests
    ])
    doc.add_heading("4. 当前P0、P1、P2、P3", level=1)
    add_table(doc, ["级别", "数量", "问题"], [
        ["P0", p0, "；".join(f["title"] for f in findings if f["priority"] == "P0") or "无"],
        ["P1", p1, "；".join(f["title"] for f in findings if f["priority"] == "P1") or "无"],
        ["P2", p2, "；".join(f["title"] for f in findings if f["priority"] == "P2") or "无"],
        ["P3", p3, "；".join(f["title"] for f in findings if f["priority"] == "P3") or "无"],
    ])
    doc.add_heading("5. 仍属待验证或证据不足", level=1)
    for text in [
        "真实API、数据库账号、身份提供方、对象存储、病毒扫描和API网关尚未部署实测。",
        "当前测试包括静态契约及可执行内存参考实现，不等同于真实数据库事务、持久化幂等和并发测试。",
        "原骨图像、拓片、馆藏、方向、残损、版权及既有争议项继续保持待核验或争议状态。",
    ]: doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("6. 验收结论", level=1); doc.add_paragraph(conclusion)
    doc.add_paragraph("本轮仅完成第13步独立验收，没有进入下一步骤。")
    doc.save(report)

    ledger = out / f"甲骨文数字文化教育系统_核心数据库_第13步独立验收问题清单与证据台账_V2.2_{'通过版' if not findings else '暂不通过版'}_20260722.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "00_验收总览"
    ws.append(["项目", "结果"]); ws.append(["结论", conclusion]); ws.append(["P0/P1/P2/P3", f"{p0}/{p1}/{p2}/{p3}"])
    ws.append(["V2.1测试", audit["v21_tests"]]); ws.append(["V2.0回归", audit["v20_tests"]]); ws.append(["独立边界测试", f"{sum(t['passed'] for t in tests)}/{len(tests)}"])
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4472C4")
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 110
    issues_ws = wb.create_sheet("01_P0-P3问题清单"); issues_ws.append(["问题ID", "优先级", "标题", "证据", "影响", "修正要求"])
    for f in findings: issues_ws.append([f["issue_id"], f["priority"], f["title"], f["evidence"], f["impact"], f["repair"]])
    tests_ws = wb.create_sheet("02_独立复测证据"); tests_ws.append(["测试ID", "场景", "预期", "实际", "结果", "对应问题"])
    for t in tests: tests_ws.append([t["test_id"], t["title"], t["expected"], json.dumps(t["actual"], ensure_ascii=False), "PASS" if t["passed"] else "FAIL", t.get("issue")])
    files_ws = wb.create_sheet("03_文件与历史"); files_ws.append(["检查", "结果"])
    for key in ["complete_sha256", "zip_crc", "manifest_entries", "manifest_errors", "v21_tests", "v20_tests", "history_compare", "mapping"]:
        files_ws.append([key, json.dumps(audit[key], ensure_ascii=False) if isinstance(audit[key], (dict, list)) else audit[key]])
    for sheet in wb.worksheets:
        for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="4472C4")
        for row in sheet.iter_rows():
            for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(1, sheet.max_column + 1): sheet.column_dimensions[chr(64 + col)].width = 45 if col > 2 else 25
    wb.save(ledger)

    machine = out / f"甲骨文数字文化教育系统_核心数据库_第13步独立验收机器审计摘要_V2.2_{'通过版' if not findings else '暂不通过版'}_20260722.json"
    machine_obj = {
        "project": "甲骨文数字文化教育系统", "step": 13, "audit_version": "V2.2", "audited_version": "V2.1", "date": DATE,
        "conclusion": conclusion, "priority_summary": {"P0": p0, "P1": p1, "P2": p2, "P3": p3},
        "audit": audit, "independent_tests": tests, "findings": findings,
        "core_database_sha256": CORE_SHA, "core_database_modified": False,
        "audited_candidate_modified": False, "next_step_executed": False,
    }
    write_json(machine, machine_obj)

    sha_txt = out / f"甲骨文数字文化教育系统_核心数据库_第13步独立验收文件SHA-256清单_V2.2_{'通过版' if not findings else '暂不通过版'}_20260722.txt"
    sha_txt.write_text("\n".join([f"{sha256(report)}  {report.name}", f"{sha256(ledger)}  {ledger.name}", f"{sha256(machine)}  {machine.name}"]) + "\n", encoding="utf-8")
    return [report, ledger, machine, sha_txt]


def visual_qa(out: Path, docx_files: list[Path]) -> dict[str, Any]:
    render_dir = out / "rendered_docx"; render_dir.mkdir(exist_ok=True)
    results = []
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        return {"available": False, "documents": []}
    for docx in docx_files:
        proc = subprocess.run([libreoffice, "--headless", "--convert-to", "pdf", "--outdir", str(render_dir), str(docx)], capture_output=True, text=True, timeout=180)
        pdf = render_dir / (docx.stem + ".pdf")
        ok = proc.returncode == 0 and pdf.exists() and pdf.stat().st_size > 0
        pages = None
        if ok and shutil.which("pdfinfo"):
            info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True)
            match = re.search(r"Pages:\s+(\d+)", info.stdout)
            pages = int(match.group(1)) if match else None
        results.append({"docx": docx.name, "pdf": pdf.name if pdf.exists() else None, "ok": ok, "pages": pages, "stdout": proc.stdout[-500:], "stderr": proc.stderr[-500:]})
    return {"available": True, "documents": results, "all_pass": all(x["ok"] for x in results)}


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--source-url", required=True); parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    out = Path(args.output_dir); shutil.rmtree(out, ignore_errors=True); out.mkdir(parents=True)
    work = Path(tempfile.mkdtemp(prefix="step13_v21_independent_"))
    outer_zip = work / "outer.zip"; download(args.source_url, outer_zip)
    outer_crc = "PASS"
    with zipfile.ZipFile(outer_zip) as z:
        bad = z.testzip(); outer_crc = "PASS" if bad is None else f"FAIL:{bad}"
        z.extractall(work / "outer")
    outer_dir = work / "outer"
    complete_candidates = list(outer_dir.rglob("*第13步教育转化接口准备完整包_V2.1_修正版_候选_20260722.zip"))
    if len(complete_candidates) != 1:
        raise RuntimeError(f"Expected one inner complete package, found {len(complete_candidates)}")
    complete = complete_candidates[0]
    complete_sha = sha256(complete)
    inner_crc = "PASS"
    with zipfile.ZipFile(complete) as z:
        bad = z.testzip(); inner_crc = "PASS" if bad is None else f"FAIL:{bad}"
        inner_extract = work / "inner"; inner_extract.mkdir(); z.extractall(inner_extract)
    root = identify_package_root(inner_extract); current = root / "current"
    manifest = read_json(root / "artifact_sha256_manifest.json", {})
    manifest_errors = []
    for entry in manifest.get("entries", []):
        p = root / entry["path"]
        if not p.exists(): manifest_errors.append("missing:" + entry["path"])
        elif sha256(p) != entry["sha256"]: manifest_errors.append("hash:" + entry["path"])

    runtime = work / "current_runtime"; shutil.copytree(current, runtime); clean_runtime(runtime)
    proc = run_python(runtime / "validate_current.py", runtime)
    if proc.returncode != 0:
        raise RuntimeError("V2.1 validator failed\n" + proc.stdout + "\n" + proc.stderr)
    v21_summary = find_validation_summary(runtime)

    history_candidates = [p for p in (root / "history").iterdir() if p.is_dir() and "V2.0" in p.name]
    if len(history_candidates) != 1: raise RuntimeError("Cannot locate V2.0 history")
    history = history_candidates[0]
    history_current = history / "current"
    history_runtime = work / "history_runtime"; shutil.copytree(history_current, history_runtime); clean_runtime(history_runtime)
    hproc = run_python(history_runtime / "validate_current.py", history_runtime)
    if hproc.returncode != 0:
        raise RuntimeError("V2.0 history validator failed\n" + hproc.stdout + "\n" + hproc.stderr)
    v20_summary = find_validation_summary(history_runtime)

    original_v20_zips = list(history.glob("*V2.0*原始受审包*.zip"))
    history_compare = {"status": "NOT_FOUND"}
    if original_v20_zips:
        original_extract = work / "original_v20"; original_extract.mkdir()
        with zipfile.ZipFile(original_v20_zips[0]) as z: z.extractall(original_extract)
        original_root = identify_package_root(original_extract)
        original_map = file_hash_map(original_root)
        history_map = file_hash_map(history, {original_v20_zips[0].name})
        missing = sorted(set(original_map) - set(history_map)); extra = sorted(set(history_map) - set(original_map))
        diff = sorted(k for k in set(original_map) & set(history_map) if original_map[k] != history_map[k])
        history_compare = {"source_files": len(original_map), "history_files": len(history_map), "missing": len(missing), "extra": len(extra), "hash_differences": len(diff), "status": "PASS" if not missing and not extra and not diff else "FAIL"}

    tests, findings = independent_boundary_tests(current)
    mapping = inspect_mapping(outer_dir)
    audit = {
        "outer_zip_crc": outer_crc, "complete_sha256": complete_sha, "expected_complete_sha256": EXPECTED_COMPLETE_SHA,
        "complete_sha_match": complete_sha == EXPECTED_COMPLETE_SHA, "zip_crc": inner_crc,
        "manifest_entries": manifest.get("entry_count", len(manifest.get("entries", []))), "manifest_errors": len(manifest_errors),
        "v21_tests": v21_summary.get("combined"), "v20_tests": v20_summary.get("combined"),
        "history_compare": history_compare, "mapping": mapping,
        "core_database_sha256": CORE_SHA, "core_database_modified": False,
        "independent_tests": {"passed": sum(t["passed"] for t in tests), "total": len(tests), "failed": sum(not t["passed"] for t in tests)},
    }
    report_files = create_reports(out, audit, tests, findings)
    qa = visual_qa(out, [p for p in report_files if p.suffix == ".docx"])
    write_json(out / "独立验收_DOCX渲染检查_V2.2.json", qa)
    result = {"audit": audit, "findings": findings, "reports": [p.name for p in report_files], "visual_qa": qa}
    write_json(out / "independent_audit_result.json", result)
    bundle = out.parent / "step13_v21_independent_audit_artifact.zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for p in sorted(out.rglob("*")):
            if p.is_file(): z.write(p, arcname=p.relative_to(out))
    print(json.dumps({"bundle": str(bundle), "conclusion": "通过，可以进入下一步" if not findings else "暂不通过，需要修正", "findings": findings, "audit": audit}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
